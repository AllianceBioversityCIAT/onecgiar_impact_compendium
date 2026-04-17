from io import BytesIO
from unittest.mock import Mock

from openpyxl import load_workbook

from app.db.connection import get_db
from app.main import app
from app.routers.reports import FULL_REPORT_COLUMNS


EXPORT_HEADERS = [
    "Study ID",
    "Title",
    "Summary",
    "Year",
    "DOI",
    "Period Start",
    "Period End",
    "Intervention Details",
    "Active",
    "Category ID",
    "Category",
    "Created At",
    "Last Updated",
    "Created By",
]
FULL_EXPORT_ROW = {
    "study_id": 101,
    "title": "Study title",
    "summary": "Summary text",
    "year": 2024,
    "period_start": "2021-01-01",
    "period_end": "2024-12-31",
    "doi": "10.1234/example",
    "category_name": "Systems",
    "intervention_details": "Intervention details",
    "countries": "Kenya, Uganda",
    "regions": "East and Southern Africa, West and Central Africa",
    "crop_types": "Maize, Wheat",
    "impact_areas_primary": "Climate Adaptation and Mitigation",
    "impact_areas_secondary": "Nutrition Health and Food Security",
    "keywords": "adaptation, resilience",
    "intervention_types": "Policy: Seed reform, Training",
    "centers": "CIMMYT, IITA",
    "initiatives": "Accelerated Breeding, Market Intelligence",
    "indicators": "Yield = 10 %; Income = 25 USD",
    "is_active": 1,
    "created_at": "2026-01-01T00:00:00",
    "last_updated_date": "2026-01-02T00:00:00",
    "created_by": "tester@example.com",
}
EMPTY_FULL_EXPORT_ROW = {
    **FULL_EXPORT_ROW,
    "study_id": 102,
    "countries": None,
    "regions": None,
    "crop_types": None,
    "impact_areas_primary": None,
    "impact_areas_secondary": None,
    "keywords": None,
    "intervention_types": None,
    "centers": None,
    "initiatives": None,
    "indicators": None,
}


class MockRow:
    def __init__(self, **kwargs):
        self.__dict__.update(kwargs)


def _export_row(active=1):
    return (
        101,
        "Study title",
        "Summary text",
        2024,
        "10.1234/example",
        "2021",
        "2024",
        "Intervention details",
        active,
        7,
        "Systems",
        "2026-01-01T00:00:00",
        "2026-01-02T00:00:00",
        "tester@example.com",
    )


def _override_db(rows):
    mock_result = Mock()
    mock_result.fetchall.return_value = rows
    mock_db = Mock()
    mock_db.execute.return_value = mock_result
    app.dependency_overrides[get_db] = lambda: mock_db
    return mock_db


def _override_full_export(rows):
    mock_session = Mock()
    mock_result = Mock()
    mock_result.fetchall.return_value = rows
    mock_session.execute.side_effect = [Mock(), mock_result]
    app.dependency_overrides[get_db] = lambda: mock_session
    return mock_session


def test_export_returns_xlsx_magic_bytes(client):
    _override_db([_export_row()])

    response = client.get("/api/reports/export?format=excel&limit=1000")

    app.dependency_overrides.clear()

    assert response.status_code == 200
    assert response.content.startswith(b"PK\x03\x04")


def test_export_content_type_and_disposition(client):
    _override_db([_export_row()])

    response = client.get("/api/reports/export?format=excel&limit=1000")

    app.dependency_overrides.clear()

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"].startswith(
        'attachment; filename="impact_compendium_full_report_'
    )
    assert response.headers["content-disposition"].endswith('.xlsx"')


def test_export_column_set(client):
    _override_db([_export_row()])

    response = client.get("/api/reports/export?format=excel&limit=1000")

    app.dependency_overrides.clear()

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook["Studies"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]

    assert headers == EXPORT_HEADERS


def test_export_404_when_no_studies(client):
    _override_db([])

    response = client.get("/api/reports/export?format=excel&limit=1000")

    app.dependency_overrides.clear()

    assert response.status_code == 404
    assert response.json() == {"detail": "No studies found for export"}


def test_export_limit_bounds(client):
    response_low = client.get("/api/reports/export?format=excel&limit=0")
    response_high = client.get("/api/reports/export?format=excel&limit=5001")

    assert response_low.status_code == 422
    assert response_high.status_code == 422


def test_export_full_returns_200_and_xlsx_magic_bytes(client):
    _override_full_export([MockRow(**FULL_EXPORT_ROW)])

    response = client.get("/api/reports/export/full?limit=5")

    app.dependency_overrides.clear()

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.content.startswith(b"PK\x03\x04")


def test_export_full_column_headers_match_spec(client):
    _override_full_export([MockRow(**FULL_EXPORT_ROW)])

    response = client.get("/api/reports/export/full?limit=5")

    app.dependency_overrides.clear()

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook["Studies (Full)"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]

    assert headers == FULL_REPORT_COLUMNS


def test_export_full_collapses_multi_value(client):
    _override_full_export([MockRow(**FULL_EXPORT_ROW)])

    response = client.get("/api/reports/export/full?limit=5")

    app.dependency_overrides.clear()

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook["Studies (Full)"]
    row = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]

    assert row[9] == "Kenya, Uganda"
    assert row[10] == "East and Southern Africa, West and Central Africa"
    assert row[12] == "Climate Adaptation and Mitigation"
    assert row[13] == "Nutrition Health and Food Security"
    assert row[18] == "Yield = 10 %; Income = 25 USD"


def test_export_full_empty_relations_render_as_empty_string(client):
    _override_full_export([MockRow(**EMPTY_FULL_EXPORT_ROW)])

    response = client.get("/api/reports/export/full?limit=5")

    app.dependency_overrides.clear()

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook["Studies (Full)"]
    row = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]

    assert row[9] in (None, "")
    assert row[10] in (None, "")
    assert row[18] in (None, "")


def test_export_full_404_when_no_studies(client):
    _override_full_export([])

    response = client.get("/api/reports/export/full?limit=5")

    app.dependency_overrides.clear()

    assert response.status_code == 404
    assert response.json() == {"detail": "No studies found for export"}
